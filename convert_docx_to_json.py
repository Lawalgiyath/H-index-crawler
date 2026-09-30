"""
Convert Word document containing staff names to JSON format
"""
import json
from docx import Document

def extract_staff_from_docx(docx_path, output_json_path, department="Chemistry"):
    """Extract staff names from Word document and save as JSON"""
    try:
        doc = Document(docx_path)
        staff_list = []
        seen_names = set()
        
        # Keywords to skip
        skip_keywords = [
            'b.sc', 'm.sc', 'ph.d', 'professor', 'department', 'chemistry',
            'analysis', 'email', 'tel:', 'fax:', 'university', 'research',
            'electroanalytical', 'medicinal', 'analytical', 'organic', 'inorganic',
            'physical', 'industrial', 'environmental'
        ]
        
        def is_likely_name(text):
            """Check if text is likely a person's name"""
            text_lower = text.lower()
            
            # Skip if contains skip keywords
            for keyword in skip_keywords:
                if keyword in text_lower:
                    return False
            
            # Skip if too short or too long
            if len(text) < 5 or len(text) > 50:
                return False
            
            # Skip if contains special characters suggesting it's not a name
            if any(char in text for char in ['@', '(', ')', ',', ';']):
                return False
            
            # Should contain at least 2 words (first and last name)
            words = text.split()
            if len(words) < 2:
                return False
            
            # Check if it looks like a name (starts with capital letters)
            if not any(word[0].isupper() for word in words if word):
                return False
            
            return True
        
        # Extract text from paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            
            # Clean up
            text = text.replace("•", "").replace("-", "").strip()
            
            if text and is_likely_name(text):
                name = ' '.join(text.split())  # Normalize whitespace
                if name not in seen_names:
                    seen_names.add(name)
                    staff_list.append({
                        "name": name,
                        "department": department
                    })
        
        # Extract from tables if any
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    text = text.replace("•", "").replace("-", "").strip()
                    
                    if text and is_likely_name(text):
                        name = ' '.join(text.split())
                        if name not in seen_names:
                            seen_names.add(name)
                            staff_list.append({
                                "name": name,
                                "department": department
                            })
        
        # Save to JSON
        with open(output_json_path, 'w', encoding='utf-8') as f:
            json.dump(staff_list, f, indent=2, ensure_ascii=False)
        
        print(f"✓ Successfully extracted {len(staff_list)} staff members")
        print(f"✓ Saved to: {output_json_path}")
        
        # Display preview
        print("\nPreview of first 5 entries:")
        for staff in staff_list[:5]:
            print(f"  - {staff['name']} ({staff['department']})")
        
        return staff_list
    
    except Exception as e:
        print(f"✗ Error: {e}")
        return []

if __name__ == "__main__":
    docx_file = "Chemistry Dept Staff.docx"
    json_file = "chemistry_staff.json"
    
    extract_staff_from_docx(docx_file, json_file, "Chemistry")
